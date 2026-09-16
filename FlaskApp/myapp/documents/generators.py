from docx import Document
from docx.shared import Inches
from flask import send_file
import io

def generate_roster_docx(board):
    doc = Document()
    doc.add_heading(f'{board.name} — Meeting Roster', level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'Name', 'Role', 'Term'

    for assignment in board.role_assignments:
        if assignment.person.honorific_title:
            honorific = assignment.person.honorific_title
        else:
            honorific = ""
        row = table.add_row().cells
        row[0].text = honorific + " " + assignment.person.first_name + ' ' + assignment.person.last_name
        row[1].text = assignment.role_type.name
        row[2].text = f'{assignment.start_date} – {assignment.end_date or "present"}'

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf